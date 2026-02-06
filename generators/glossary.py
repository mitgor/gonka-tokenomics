"""
Gonka Tokenomics - Glossary / Definitions Tab Builder

Builds a "Definitions" worksheet with alphabetically sorted term/definition
pairs for standalone workbooks (REQ-U11).

Public API:
    build_glossary_tab(wb, terms) -> None
"""

from openpyxl.styles import Font, Alignment

from generators.styles import TAB_COLOR_DOCS


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def build_glossary_tab(wb, terms):
    """Build a Definitions/Glossary tab with alphabetical entries.

    Args:
        wb: openpyxl Workbook (styles already registered via register_styles).
        terms: dict of {term: definition} to display.

    Returns:
        None -- modifies ``wb`` in-place.
    """
    ws = wb.create_sheet("Definitions")
    ws.sheet_properties.tabColor = TAB_COLOR_DOCS

    # =================================================================
    # Title row (row 1)
    # =================================================================
    ws.merge_cells("A1:C1")
    ws.cell(row=1, column=1, value="DEFINITIONS / GLOSSARY").style = "section_header"

    # Blank separator row 2

    # =================================================================
    # Column headers (row 3)
    # =================================================================
    ws.cell(row=3, column=1, value="Term").style = "header"
    ws.cell(row=3, column=2, value="Definition").style = "header"

    # =================================================================
    # Data rows (row 4+) -- alphabetically sorted
    # =================================================================
    row = 4
    for term in sorted(terms.keys()):
        definition = terms[term]

        ws.cell(row=row, column=1, value=term).font = Font(
            name="Calibri", size=11, bold=True,
        )

        def_cell = ws.cell(row=row, column=2, value=definition)
        def_cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Row height heuristic: accommodate wrapped text
        ws.row_dimensions[row].height = max(15, len(definition) // 80 * 15 + 15)

        row += 1

    # =================================================================
    # Column widths
    # =================================================================
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 15

    # =================================================================
    # Freeze panes at A4 (header row always visible)
    # =================================================================
    ws.freeze_panes = "A4"
