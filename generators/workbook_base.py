"""
Gonka Tokenomics - Workbook Base (Assumptions Tab Builder)

Core function: build_assumptions_tab(wb) writes all parameters from PARAM_GROUPS
to a formatted Assumptions sheet and returns the param_refs dict.

build_filtered_assumptions_tab(wb, param_names) writes only a subset of
parameters and returns correctly-addressed param_refs, enabling standalone
workbooks to reuse the same model builders with remapped cell addresses.

The param_refs dict is THE critical interface contract: every downstream model
module (emission, token_price, fee_transition, etc.) receives this dict to
write Excel formulas referencing Assumptions!$B$N cells.

create_workbook() is the convenience entry point that creates a fresh workbook,
builds the Assumptions tab, and returns (wb, param_refs).
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from models.parameters import PARAM_GROUPS
from generators.styles import (
    register_styles,
    FORMAT_TO_STYLE,
    INPUT_FILL_COLOR,
    INPUT_FONT_COLOR,
    TAB_COLOR_INPUT,
)


# Column widths for the Assumptions tab
_COL_WIDTHS = {
    "A": 35,
    "B": 20,
    "C": 15,
    "D": 40,   # Wider for full document-reference citations
    "E": 12,   # Confidence column (HIGH/MED/LOW)
}

# Font for source citations (italic)
_SOURCE_FONT = Font(name="Calibri", size=11, italic=True)


def _add_tail_emission_toggle(ws, param_refs):
    """Add ON/OFF DataValidation to the Tail Emission Toggle parameter cell.

    The cell already exists from the PARAM_GROUPS loop. This function
    adds a dropdown constraint so the user can only enter ON or OFF.

    Args:
        ws: The Assumptions worksheet.
        param_refs: Dict containing "Tail Emission Toggle" -> "Assumptions!$B$NN".
    """
    # Parse row number from param_refs (format: "Assumptions!$B$NN")
    ref = param_refs["Tail Emission Toggle"]
    row = int(ref.split("$")[-1])

    toggle_cell = ws.cell(row=row, column=2)

    dv = DataValidation(
        type="list",
        formula1='"ON,OFF"',
        allow_blank=False,
    )
    dv.prompt = "Select ON or OFF"
    dv.promptTitle = "Tail Emission"
    ws.add_data_validation(dv)
    dv.add(toggle_cell)


def _add_scenario_selector(ws, current_row, param_refs):
    """Add scenario selector dropdown and CHOOSE formulas to Assumptions tab.

    Adds a SCENARIO SELECTOR section with:
      - DataValidation dropdown (Conservative/Base/Aggressive)
      - MATCH formula converting text to index 1/2/3
      - CHOOSE formulas for Active Price Low and Active Price High

    Args:
        ws: The Assumptions worksheet.
        current_row: Next available row after PARAM_GROUPS.
        param_refs: Dict to update with new scenario cell references.

    Returns:
        int: The next available row after the scenario selector section.
    """
    # Section header
    ws.merge_cells(
        start_row=current_row, start_column=1,
        end_row=current_row, end_column=5,
    )
    section_cell = ws.cell(row=current_row, column=1, value="SCENARIO SELECTOR")
    section_cell.style = "section_header"
    current_row += 1

    # --- Active Scenario dropdown ---
    ws.cell(row=current_row, column=1, value="Active Scenario")
    selector_cell = ws.cell(row=current_row, column=2, value="Base")
    selector_cell.style = "input_cell"
    selector_cell.protection = Protection(locked=False)

    dv = DataValidation(
        type="list",
        formula1='"Conservative,Base,Aggressive"',
        allow_blank=False,
    )
    dv.prompt = "Select scenario"
    dv.promptTitle = "Active Scenario"
    ws.add_data_validation(dv)
    dv.add(selector_cell)

    param_refs["Active Scenario"] = f"Assumptions!$B${current_row}"
    selector_row = current_row
    current_row += 1

    # --- Scenario Index (MATCH converts text -> 1/2/3) ---
    ws.cell(row=current_row, column=1, value="Scenario Index")
    index_cell = ws.cell(
        row=current_row, column=2,
        value=f'=MATCH(B{selector_row},{{"Conservative","Base","Aggressive"}},0)',
    )
    index_cell.style = "formula_cell"

    param_refs["Scenario Index"] = f"Assumptions!$B${current_row}"
    index_row = current_row
    current_row += 1

    # --- Active Price Low (CHOOSE selects per scenario) ---
    ws.cell(row=current_row, column=1, value="Active Price Low")
    low_cell = ws.cell(
        row=current_row, column=2,
        value=(
            f"=CHOOSE($B${index_row},"
            f"{param_refs['Conservative Price Low']},"
            f"{param_refs['Moderate Price Low']},"
            f"{param_refs['Aggressive Price Low']})"
        ),
    )
    low_cell.style = "formula_cell"
    low_cell.number_format = "$#,##0.00"

    param_refs["Active Price Low"] = f"Assumptions!$B${current_row}"
    current_row += 1

    # --- Active Price High (CHOOSE selects per scenario) ---
    ws.cell(row=current_row, column=1, value="Active Price High")
    high_cell = ws.cell(
        row=current_row, column=2,
        value=(
            f"=CHOOSE($B${index_row},"
            f"{param_refs['Conservative Price High']},"
            f"{param_refs['Moderate Price High']},"
            f"{param_refs['Aggressive Price High']})"
        ),
    )
    high_cell.style = "formula_cell"
    high_cell.number_format = "$#,##0.00"

    param_refs["Active Price High"] = f"Assumptions!$B${current_row}"
    current_row += 1

    return current_row


def _add_what_if_toggle(ws, param_refs, toggle_name):
    """Add Y/N DataValidation to a what-if toggle parameter cell.

    The cell already exists from the PARAM_GROUPS loop. This function
    adds a dropdown constraint so the user can only enter Y or N.

    Args:
        ws: The Assumptions worksheet.
        param_refs: Dict containing toggle_name -> "Assumptions!$B$NN".
        toggle_name: The parameter name (e.g., "Buyback-Burn Active").
    """
    ref = param_refs[toggle_name]
    row = int(ref.split("$")[-1])

    toggle_cell = ws.cell(row=row, column=2)

    dv = DataValidation(
        type="list",
        formula1='"Y,N"',
        allow_blank=False,
    )
    dv.prompt = "Select Y or N"
    dv.promptTitle = toggle_name
    ws.add_data_validation(dv)
    dv.add(toggle_cell)


def build_filtered_assumptions_tab(wb, param_names,
                                   include_scenario_selector=True,
                                   include_tail_toggle=False,
                                   include_toggles=None):
    """Build Assumptions tab with only specified parameters.

    Modeled on build_assumptions_tab() but writes only parameters whose
    names are in param_names. Returns a param_refs dict with correctly
    remapped cell addresses so existing model builders work unchanged.

    Args:
        wb: An openpyxl Workbook instance.
        param_names: Set/frozenset of parameter name strings to include.
        include_scenario_selector: Whether to add the scenario selector
            section (requires 6 price scenario params). Default True.
        include_tail_toggle: Whether to add ON/OFF DataValidation to
            Tail Emission Toggle cell. Default False.
        include_toggles: List of what-if toggle names to add Y/N
            DataValidation to (e.g., ["Buyback-Burn Active"]).
            Default None (no toggles).

    Returns:
        dict: Mapping of parameter name -> "Assumptions!$B$N" cell reference.
    """
    # Register all named styles with this workbook
    register_styles(wb)

    # Get or create the active worksheet, rename to "Assumptions"
    ws = wb.active
    ws.title = "Assumptions"
    ws.sheet_properties.tabColor = TAB_COLOR_INPUT

    # --- Row 1: Main header (merged A1:F1) ---
    ws.merge_cells("A1:F1")
    cell_a1 = ws["A1"]
    cell_a1.value = "GONKA TOKENOMICS MODEL - ASSUMPTIONS"
    cell_a1.style = "section_header"

    # --- Row 2: Instruction text (merged A2:F2) ---
    ws.merge_cells("A2:F2")
    cell_a2 = ws["A2"]
    cell_a2.value = "All blue-shaded cells below are adjustable inputs"

    # --- Row 3: Blank separator ---
    current_row = 4  # Start data from row 4

    # --- Build param_refs dict as we write filtered parameters ---
    param_refs = {}

    for group_name, params in PARAM_GROUPS.items():
        # Filter: only include params whose names are in param_names
        filtered_params = [p for p in params if p["name"] in param_names]
        if not filtered_params:
            continue

        # Section header row
        ws.merge_cells(
            start_row=current_row, start_column=1,
            end_row=current_row, end_column=5,
        )
        section_cell = ws.cell(row=current_row, column=1)
        section_cell.value = group_name
        section_cell.style = "section_header"
        current_row += 1

        # Parameter rows
        for param in filtered_params:
            # Column A: Parameter name (locked by default)
            ws.cell(row=current_row, column=1, value=param["name"])

            # Column B: Parameter value (UNLOCKED input cell)
            value_cell = ws.cell(row=current_row, column=2, value=param["value"])
            value_cell.style = "input_cell"

            # Override number_format based on parameter's format type
            format_key = param["format"]
            if format_key in FORMAT_TO_STYLE:
                style_name = FORMAT_TO_STYLE[format_key]
                if style_name is not None:
                    named_style = wb._named_styles[style_name]
                    value_cell.number_format = named_style.number_format
            value_cell.protection = Protection(locked=False)

            # Column C: Unit string (locked by default)
            ws.cell(row=current_row, column=3, value=param["unit"])

            # Column D: Source citation (locked, italic)
            source_cell = ws.cell(row=current_row, column=4, value=param["source"])
            source_cell.font = _SOURCE_FONT

            # Column E: Confidence level (locked, italic)
            confidence_cell = ws.cell(row=current_row, column=5, value=param.get("confidence", ""))
            confidence_cell.font = _SOURCE_FONT

            # Record the cell reference for this parameter
            param_refs[param["name"]] = f"Assumptions!$B${current_row}"

            current_row += 1

        # Blank separator row after each group
        current_row += 1

    # --- Tail Emission Toggle (DataValidation dropdown) ---
    if include_tail_toggle and "Tail Emission Toggle" in param_refs:
        _add_tail_emission_toggle(ws, param_refs)

    # --- What-If Toggles (Y/N DataValidation dropdowns) ---
    if include_toggles:
        for toggle_name in include_toggles:
            if toggle_name in param_refs:
                _add_what_if_toggle(ws, param_refs, toggle_name)

    # --- Scenario Selector (dropdown + MATCH + CHOOSE) ---
    if include_scenario_selector:
        # Verify all 6 price scenario params are present
        scenario_params = {
            "Conservative Price Low", "Conservative Price High",
            "Moderate Price Low", "Moderate Price High",
            "Aggressive Price Low", "Aggressive Price High",
        }
        if scenario_params.issubset(param_refs.keys()):
            current_row = _add_scenario_selector(ws, current_row, param_refs)
        else:
            missing = scenario_params - param_refs.keys()
            import warnings
            warnings.warn(
                f"Scenario selector skipped: missing params {missing}"
            )

    # --- Set column widths ---
    for col_letter, width in _COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # --- Freeze panes at A4 (header rows always visible) ---
    ws.freeze_panes = "A4"

    return param_refs


def build_assumptions_tab(wb):
    """Build the Assumptions tab with all parameters from PARAM_GROUPS.

    Takes an openpyxl Workbook (already created), writes the formatted
    Assumptions sheet, and returns a param_refs dict mapping every parameter
    name to its absolute cell reference (e.g., "Assumptions!$B$12").

    Args:
        wb: An openpyxl Workbook instance.

    Returns:
        dict: Mapping of parameter name -> "Assumptions!$B$N" cell reference.
    """
    # Register all named styles with this workbook
    register_styles(wb)

    # Get or create the active worksheet, rename to "Assumptions"
    ws = wb.active
    ws.title = "Assumptions"
    ws.sheet_properties.tabColor = TAB_COLOR_INPUT

    # --- Row 1: Main header (merged A1:F1) ---
    ws.merge_cells("A1:F1")
    cell_a1 = ws["A1"]
    cell_a1.value = "GONKA TOKENOMICS MODEL - ASSUMPTIONS"
    cell_a1.style = "section_header"

    # --- Row 2: Instruction text (merged A2:F2) ---
    ws.merge_cells("A2:F2")
    cell_a2 = ws["A2"]
    cell_a2.value = "All blue-shaded cells below are adjustable inputs"

    # --- Row 3: Blank separator ---
    current_row = 4  # Start data from row 4

    # --- Build param_refs dict as we write parameters ---
    param_refs = {}

    for group_name, params in PARAM_GROUPS.items():
        # Section header row
        ws.merge_cells(
            start_row=current_row, start_column=1,
            end_row=current_row, end_column=5,
        )
        section_cell = ws.cell(row=current_row, column=1)
        section_cell.value = group_name
        section_cell.style = "section_header"
        current_row += 1

        # Parameter rows
        for param in params:
            # Column A: Parameter name (locked by default)
            ws.cell(row=current_row, column=1, value=param["name"])

            # Column B: Parameter value (UNLOCKED input cell)
            value_cell = ws.cell(row=current_row, column=2, value=param["value"])
            value_cell.style = "input_cell"

            # Override number_format based on parameter's format type
            format_key = param["format"]
            if format_key in FORMAT_TO_STYLE:
                style_name = FORMAT_TO_STYLE[format_key]
                if style_name is not None:
                    # Look up the number_format from the registered named style
                    named_style = wb._named_styles[style_name]
                    value_cell.number_format = named_style.number_format
            value_cell.protection = Protection(locked=False)

            # Column C: Unit string (locked by default)
            ws.cell(row=current_row, column=3, value=param["unit"])

            # Column D: Source citation (locked, italic)
            source_cell = ws.cell(row=current_row, column=4, value=param["source"])
            source_cell.font = _SOURCE_FONT

            # Column E: Confidence level (locked, italic)
            confidence_cell = ws.cell(row=current_row, column=5, value=param.get("confidence", ""))
            confidence_cell.font = _SOURCE_FONT

            # Record the cell reference for this parameter
            param_refs[param["name"]] = f"Assumptions!$B${current_row}"

            current_row += 1

        # Blank separator row after each group
        current_row += 1

    # --- Tail Emission Toggle (DataValidation dropdown) ---
    _add_tail_emission_toggle(ws, param_refs)

    # --- What-If Toggles (Y/N DataValidation dropdowns) ---
    for toggle_name in ("Buyback-Burn Active", "Deploy POL Active"):
        if toggle_name in param_refs:
            _add_what_if_toggle(ws, param_refs, toggle_name)

    # --- Scenario Selector (dropdown + MATCH + CHOOSE) ---
    current_row = _add_scenario_selector(ws, current_row, param_refs)

    # --- Set column widths ---
    for col_letter, width in _COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # --- Freeze panes at A4 (header rows always visible) ---
    ws.freeze_panes = "A4"

    return param_refs


def create_workbook():
    """Create a new workbook with styles registered and Assumptions tab built.

    Returns:
        tuple: (wb, param_refs) where wb is the openpyxl Workbook and
               param_refs maps parameter names to Assumptions!$B$N addresses.
    """
    wb = Workbook()
    param_refs = build_assumptions_tab(wb)
    return wb, param_refs
