"""
Gonka Tokenomics - Treasury & POL Simulation Tab

Builds the "Treasury & POL" worksheet with 32 data rows (matching the
Emission Schedule period structure) and 14 columns:

  A  Period Label               (cross-ref from Emission Schedule)
  B  CP Balance (GNK)           (waterfall: 120M - POL - defense draws)
  C  CP Outflows (GNK)          (period outflows for waterfall visibility)
  D  POL Revenue ($)            (midpoint LP fee revenue, prorated, minus rebalancing)
  E  Cumul. POL Revenue ($)     (running sum of D)
  F  Buyback Burn (GNK)         (cross-ref Token Price column K)
  G  Cumul. Burn (GNK)          (running sum of F)
  H  Burn % of Supply           (cumulative burn / total supply)
  I  AI Fund Inflows ($)        (cross-ref Fee Transition column L)
  J  AI Fund Balance ($)        (running balance: inflows - expenses)
  K  Defense GNK Allocated      (annual 6M GNK prorated)
  L  Defense Treasury ($)       (GNK converted to USD at active price, accumulated)
  M  Net Treasury (GNK)         (CP balance + POL GNK allocation)
  N  Net Treasury ($)           (all assets valued in USD)

Plus three below-data analysis sections:
  - TIME-TO-X MILESTONES (rows 36-39): depletion/target callouts
  - FLOOR DEFENSE SCENARIOS (rows 41-47): 3/6/12 month spending scenarios
  - IL CAVEAT (row 49): impermanent loss deferral notice

Every formula references the Assumptions tab via param_refs or the Emission
Schedule / Token Price / Fee Transition tabs via their meta dicts, so changes
propagate automatically.

Public API:
    build_treasury_tab(wb, param_refs, emission_meta, price_meta, fee_meta) -> treasury_meta dict
"""

from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import quote_sheetname

from generators.chart_utils import col_to_num
from generators.styles import TAB_COLOR_CALC


# ---------------------------------------------------------------------------
# Column layout
# ---------------------------------------------------------------------------

_HEADERS = [
    "Period",                     # A
    "CP Balance (GNK)",           # B
    "CP Outflows (GNK)",          # C
    "POL Revenue ($)",            # D
    "Cumul. POL Revenue ($)",     # E
    "Buyback Burn (GNK)",         # F
    "Cumul. Burn (GNK)",          # G
    "Burn % of Supply",           # H
    "AI Fund Inflows ($)",        # I
    "AI Fund Balance ($)",        # J
    "Defense GNK Allocated",      # K
    "Defense Treasury ($)",       # L
    "Net Treasury (GNK)",         # M
    "Net Treasury ($)",           # N
]

_COL_WIDTHS = {
    "A": 12, "B": 18, "C": 18, "D": 18, "E": 20,
    "F": 18, "G": 18, "H": 14, "I": 18, "J": 18,
    "K": 20, "L": 18, "M": 18, "N": 18,
}


# ---------------------------------------------------------------------------
# Below-data sections (private helpers)
# ---------------------------------------------------------------------------

def _build_time_to_x_callouts(ws, treasury_meta, defense_target_low_ref):
    """Build Time-to-X milestone callouts at rows 36-39."""
    # Row 36: Section header
    ws.merge_cells("A36:F36")
    header_cell = ws.cell(row=36, column=1, value="TIME-TO-X MILESTONES")
    header_cell.style = "section_header"

    # Row 37: Community Pool depletion
    ws.cell(row=37, column=1, value="Community Pool depleted in:")
    ws.cell(row=37, column=1).font = Font(name="Calibri", size=11, bold=True)
    formula_37 = (
        '=IFERROR(INDEX($A$3:$A$34,MATCH(TRUE,INDEX($B$3:$B$34<=0,0),0)),'
        '"Not depleted within 10 years")'
    )
    ws.cell(row=37, column=2, value=formula_37).style = "formula_cell"

    # Row 38: Floor defense reaches $2M target
    ws.cell(row=38, column=1, value="Floor defense reaches $2M target:")
    ws.cell(row=38, column=1).font = Font(name="Calibri", size=11, bold=True)
    formula_38 = (
        f'=IFERROR(INDEX($A$3:$A$34,MATCH(TRUE,INDEX($L$3:$L$34>={defense_target_low_ref},0),0)),'
        f'"Not reached within 10 years")'
    )
    ws.cell(row=38, column=2, value=formula_38).style = "formula_cell"

    # Row 39: Buyback burns 1% of supply
    ws.cell(row=39, column=1, value="Buyback burns 1% of supply:")
    ws.cell(row=39, column=1).font = Font(name="Calibri", size=11, bold=True)
    formula_39 = (
        '=IFERROR(INDEX($A$3:$A$34,MATCH(TRUE,INDEX($H$3:$H$34>=0.01,0),0)),'
        '"Not reached within 10 years")'
    )
    ws.cell(row=39, column=2, value=formula_39).style = "formula_cell"


def _build_defense_scenario_table(ws, defense_target_low_ref):
    """Build floor defense spending scenario table at rows 41-47."""
    # Row 41: Section header
    ws.merge_cells("A41:F41")
    header_cell = ws.cell(row=41, column=1, value="FLOOR DEFENSE SCENARIOS")
    header_cell.style = "section_header"

    # Row 42: Sub-headers
    sub_headers = [
        "Defense Duration",
        "Monthly Spend Rate",
        "Total Budget ($)",
        "Remaining Treasury ($)",
        "Defense Adequate?",
    ]
    for col_idx, text in enumerate(sub_headers, start=1):
        cell = ws.cell(row=42, column=col_idx, value=text)
        cell.style = "header"

    # Rows 43-45: Three duration scenarios (3, 6, 12 months)
    # Uses Defense Treasury balance at end of Year 2 (row 26 = period index 23,
    # last monthly period) as the baseline.
    durations = [3, 6, 12]
    for d_idx, n_months in enumerate(durations):
        row = 43 + d_idx

        # A: Duration label
        ws.cell(row=row, column=1, value=f"{n_months} months")

        # B: Monthly spend rate = $2M target / N months
        ws.cell(
            row=row, column=2,
            value=f"={defense_target_low_ref}/{n_months}",
        ).style = "currency"

        # C: Total budget (from low target)
        ws.cell(
            row=row, column=3,
            value=f"={defense_target_low_ref}",
        ).style = "currency"

        # D: Remaining defense treasury after full spend (Year 2 balance as baseline)
        ws.cell(
            row=row, column=4,
            value=f"=MAX(0,L26-{defense_target_low_ref})",
        ).style = "currency"

        # E: Adequacy check
        ws.cell(
            row=row, column=5,
            value=f'=IF(L26>={defense_target_low_ref},"Adequate","Insufficient")',
        )

    # Row 47: Annotation
    note_cell = ws.cell(
        row=47, column=1,
        value="* Uses Defense Treasury balance at end of Year 2 as baseline. "
              "Active defense spending modeled against $2M target.",
    )
    note_cell.font = Font(name="Calibri", size=10, italic=True)


def _build_il_caveat(ws):
    """Build IL caveat label at row 49."""
    ws.merge_cells("A49:F49")
    caveat_cell = ws.cell(
        row=49, column=1,
        value="IL impact not modeled; see v2 for concentrated position risk analysis",
    )
    caveat_cell.font = Font(name="Calibri", size=11, italic=True, color="9C0006")


# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------

def _create_treasury_composition_chart(ws, meta):
    """Stacked area: USD treasury components + Net Treasury total line."""
    chart = AreaChart()
    chart.grouping = "stacked"
    chart.title = "Treasury Components (USD) Over Time"
    chart.y_axis.title = "USD Value"
    chart.x_axis.title = "Period"
    chart.style = 13
    chart.width = 20
    chart.height = 12

    # Series: AI Fund Balance (J), Cumul POL Revenue (E), Defense Treasury (L)
    for col_key in ["ai_fund_balance", "cumul_pol_revenue", "defense_treasury"]:
        col_num = col_to_num(meta["cols"][col_key])
        data = Reference(ws, min_col=col_num, min_row=meta["header_row"],
                         max_row=meta["data_end_row"])
        chart.add_data(data, titles_from_data=True)

    # Categories: Period labels (column A)
    cats = Reference(ws, min_col=1, min_row=meta["data_start_row"],
                     max_row=meta["data_end_row"])
    chart.set_categories(cats)

    ws.add_chart(chart, "P1")


def _create_depletion_chart(ws, meta):
    """Line chart: CP balance declining + defense treasury accumulating."""
    chart = LineChart()
    chart.title = "Community Pool Depletion & Defense Treasury Growth"
    chart.y_axis.title = "Value"
    chart.x_axis.title = "Period"
    chart.style = 13
    chart.width = 20
    chart.height = 12

    # Series 1: CP Balance GNK (column B)
    cp_col = col_to_num(meta["cols"]["cp_balance"])
    data1 = Reference(ws, min_col=cp_col, min_row=meta["header_row"],
                      max_row=meta["data_end_row"])
    chart.add_data(data1, titles_from_data=True)

    # Series 2: Defense Treasury USD (column L)
    def_col = col_to_num(meta["cols"]["defense_treasury"])
    data2 = Reference(ws, min_col=def_col, min_row=meta["header_row"],
                      max_row=meta["data_end_row"])
    chart.add_data(data2, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=meta["data_start_row"],
                     max_row=meta["data_end_row"])
    chart.set_categories(cats)

    ws.add_chart(chart, "P17")


def _create_buyback_burn_chart(ws, meta):
    """Bar chart: cumulative buyback burn with % of supply overlay."""
    chart = BarChart()
    chart.type = "col"
    chart.title = "Cumulative Buyback-Burn (GNK)"
    chart.y_axis.title = "GNK Burned"
    chart.x_axis.title = "Period"
    chart.style = 13
    chart.width = 20
    chart.height = 12

    # Series 1: Cumulative Burn GNK (column G) - bars
    burn_col = col_to_num(meta["cols"]["cumul_burn"])
    data1 = Reference(ws, min_col=burn_col, min_row=meta["header_row"],
                      max_row=meta["data_end_row"])
    chart.add_data(data1, titles_from_data=True)

    # Series 2: Burn % of Supply (column H) - line on secondary axis
    pct_col = col_to_num(meta["cols"]["burn_pct_supply"])
    data2 = Reference(ws, min_col=pct_col, min_row=meta["header_row"],
                      max_row=meta["data_end_row"])

    # Create secondary line chart
    line = LineChart()
    line.add_data(data2, titles_from_data=True)
    line.y_axis.title = "% of Total Supply"
    line.y_axis.axId = 200

    # Categories
    cats = Reference(ws, min_col=1, min_row=meta["data_start_row"],
                     max_row=meta["data_end_row"])
    chart.set_categories(cats)
    line.set_categories(cats)

    # Style the line series
    s = line.series[0]
    s.graphicalProperties.line.width = 25000

    chart.y_axis.crosses = "min"
    chart += line

    ws.add_chart(chart, "P33")


# ---------------------------------------------------------------------------
# Conditional Formatting
# ---------------------------------------------------------------------------

def _add_cp_balance_formatting(ws, meta):
    """Red when CP < 10M GNK, green when > 50M GNK."""
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(name="Calibri", size=11, color="9C0006")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    green_font = Font(name="Calibri", size=11, color="006100")

    cell_range = f"B{meta['data_start_row']}:B{meta['data_end_row']}"

    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(operator="greaterThan", formula=["50000000"],
                   fill=green_fill, font=green_font),
    )
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(operator="lessThan", formula=["10000000"],
                   fill=red_fill, font=red_font),
    )


def _add_defense_treasury_formatting(ws, meta):
    """Red when defense < $1M, green when > $3M."""
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(name="Calibri", size=11, color="9C0006")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    green_font = Font(name="Calibri", size=11, color="006100")

    cell_range = f"L{meta['data_start_row']}:L{meta['data_end_row']}"

    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(operator="greaterThan", formula=["3000000"],
                   fill=green_fill, font=green_font),
    )
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(operator="lessThan", formula=["1000000"],
                   fill=red_fill, font=red_font),
    )


def _add_net_treasury_heatmap(ws, meta):
    """ColorScaleRule: red (low) -> yellow (mid) -> green (high) on net treasury."""
    cell_range = f"N{meta['data_start_row']}:N{meta['data_end_row']}"

    rule = ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B",          # Red
        mid_type="num", mid_value=50000000, mid_color="FFEB84",         # Yellow ($50M)
        end_type="num", end_value=200000000, end_color="63BE7B",        # Green ($200M)
    )
    ws.conditional_formatting.add(cell_range, rule)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def build_treasury_tab(wb, param_refs, emission_meta, price_meta, fee_meta):
    """Create the 'Treasury & POL' worksheet and populate formulas.

    Args:
        wb: An openpyxl Workbook with styles already registered.
        param_refs: dict mapping parameter names to "Assumptions!$B$N".
        emission_meta: dict from build_emission_tab() with sheet coordinates.
        price_meta: dict from build_token_price_tab() with sheet coordinates.
        fee_meta: dict from build_fee_transition_tab() with sheet coordinates.

    Returns:
        dict: treasury_meta with sheet coordinates for downstream tabs/charts.
    """
    ws = wb.create_sheet(title="Treasury & POL")
    ws.sheet_properties.tabColor = TAB_COLOR_CALC

    # ------------------------------------------------------------------
    # Resolve param_refs
    # ------------------------------------------------------------------
    cp_ref = param_refs["Community Pool"]
    pol_alloc_ref = param_refs["POL GNK Allocation"]
    total_supply_ref = param_refs["Total Supply"]
    pol_rev_low_ref = param_refs["Expected LP Fee Revenue Low"]
    pol_rev_high_ref = param_refs["Expected LP Fee Revenue High"]
    pol_rebal_ref = param_refs["POL Rebalancing Cost"]
    defense_annual_ref = param_refs["Annual GNK Allocation for Defense"]
    defense_target_low_ref = param_refs["Defense Treasury Target Low"]
    ai_fund_expenses_ref = param_refs["AI Fund Monthly Expenses"]

    # ------------------------------------------------------------------
    # Cross-sheet references
    # ------------------------------------------------------------------
    es_sheet = quote_sheetname(emission_meta["sheet_name"])
    tp_sheet = quote_sheetname(price_meta["sheet_name"])
    ft_sheet = quote_sheetname(fee_meta["sheet_name"])

    es_period_col = emission_meta["cols"]["period_label"]       # "A"
    tp_active_col = price_meta["cols"]["active_price"]          # "F"
    tp_buyback_col = price_meta["cols"]["buyback_burn"]         # "K"
    ft_ai_fund_col = fee_meta["cols"]["ai_fund_share"]          # "L"

    # ------------------------------------------------------------------
    # Row 1: Merged title
    # ------------------------------------------------------------------
    ws.merge_cells("A1:N1")
    title_cell = ws.cell(row=1, column=1, value="TREASURY & POL SIMULATION")
    title_cell.style = "section_header"

    # ------------------------------------------------------------------
    # Row 2: Column headers
    # ------------------------------------------------------------------
    for col_idx, header_text in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header_text)
        cell.style = "header"

    # ------------------------------------------------------------------
    # Rows 3-34: Data rows (32 periods)
    # ------------------------------------------------------------------
    data_start_row = 3
    num_periods = 32

    for i in range(num_periods):
        row = data_start_row + i
        days = 30 if i < 24 else 365

        # Cross-sheet row references
        es_row = emission_meta["data_start_row"] + i
        tp_row = price_meta["data_start_row"] + i
        ft_row = fee_meta["data_start_row"] + i

        # A: Period Label (cross-ref from Emission Schedule)
        ws.cell(
            row=row, column=1,
            value=f"={es_sheet}!{es_period_col}{es_row}",
        ).style = "crossref_cell"

        # B: CP Balance (GNK) -- running waterfall balance
        if i == 0:
            # Period 0: Starting balance minus POL (one-time) and defense draw
            b_formula = (
                f"={cp_ref}-{pol_alloc_ref}"
                f"-{defense_annual_ref}*{days}/365"
            )
        else:
            # Subsequent: prior balance minus defense draw
            b_formula = f"=B{row - 1}-{defense_annual_ref}*{days}/365"
        ws.cell(row=row, column=2, value=b_formula).style = "tokens"

        # C: CP Outflows (GNK) -- period outflows for waterfall visibility
        if i == 0:
            c_formula = f"={pol_alloc_ref}+{defense_annual_ref}*{days}/365"
        else:
            c_formula = f"={defense_annual_ref}*{days}/365"
        ws.cell(row=row, column=3, value=c_formula).style = "tokens"

        # D: POL Revenue ($) -- midpoint LP fee revenue minus rebalancing cost
        d_formula = (
            f"=({pol_rev_low_ref}+{pol_rev_high_ref})/2*{days}/365"
            f"-{pol_rebal_ref}*{days}/365"
        )
        ws.cell(row=row, column=4, value=d_formula).style = "currency"

        # E: Cumulative POL Revenue ($) -- running sum
        if i == 0:
            e_formula = f"=D{row}"
        else:
            e_formula = f"=E{row - 1}+D{row}"
        ws.cell(row=row, column=5, value=e_formula).style = "currency"

        # F: Buyback Burn (GNK) -- cross-ref Token Price column K
        f_formula = f"={tp_sheet}!{tp_buyback_col}{tp_row}"
        ws.cell(row=row, column=6, value=f_formula).style = "tokens"

        # G: Cumulative Burn (GNK) -- running total
        if i == 0:
            g_formula = f"=F{row}"
        else:
            g_formula = f"=G{row - 1}+F{row}"
        ws.cell(row=row, column=7, value=g_formula).style = "tokens"

        # H: Burn % of Supply
        h_formula = f"=G{row}/{total_supply_ref}"
        ws.cell(row=row, column=8, value=h_formula).style = "percent"

        # I: AI Fund Inflows ($) -- cross-ref Fee Transition column L
        i_formula = f"={ft_sheet}!{ft_ai_fund_col}{ft_row}"
        ws.cell(row=row, column=9, value=i_formula).style = "currency"

        # J: AI Fund Balance ($) -- running balance: inflows minus expenses
        if i == 0:
            j_formula = f"=I{row}-{ai_fund_expenses_ref}*{days}/30"
        else:
            j_formula = f"=J{row - 1}+I{row}-{ai_fund_expenses_ref}*{days}/30"
        ws.cell(row=row, column=10, value=j_formula).style = "currency"

        # K: Defense GNK Allocated -- annual 6M GNK prorated
        k_formula = f"={defense_annual_ref}*{days}/365"
        ws.cell(row=row, column=11, value=k_formula).style = "tokens"

        # L: Defense Treasury ($) -- GNK converted to USD, accumulated
        if i == 0:
            l_formula = f"=K{row}*{tp_sheet}!{tp_active_col}{tp_row}"
        else:
            l_formula = (
                f"=L{row - 1}+K{row}*{tp_sheet}!{tp_active_col}{tp_row}"
            )
        ws.cell(row=row, column=12, value=l_formula).style = "currency"

        # M: Net Treasury (GNK) -- CP Balance + POL GNK allocation
        m_formula = f"=B{row}+{pol_alloc_ref}"
        ws.cell(row=row, column=13, value=m_formula).style = "tokens"

        # N: Net Treasury ($) -- all assets in USD
        n_formula = (
            f"=B{row}*{tp_sheet}!{tp_active_col}{tp_row}"
            f"+{pol_alloc_ref}*{tp_sheet}!{tp_active_col}{tp_row}"
            f"+E{row}+L{row}+J{row}"
        )
        ws.cell(row=row, column=14, value=n_formula).style = "currency"

    # ------------------------------------------------------------------
    # Below-data sections
    # ------------------------------------------------------------------
    treasury_meta = {
        "sheet_name": "Treasury & POL",
        "header_row": 2,
        "data_start_row": 3,
        "data_end_row": 34,
        "cols": {
            "period_label": "A",
            "cp_balance": "B",
            "cp_outflows": "C",
            "pol_revenue": "D",
            "cumul_pol_revenue": "E",
            "buyback_burn": "F",
            "cumul_burn": "G",
            "burn_pct_supply": "H",
            "ai_fund_inflows": "I",
            "ai_fund_balance": "J",
            "defense_gnk_allocated": "K",
            "defense_treasury": "L",
            "net_treasury_gnk": "M",
            "net_treasury_usd": "N",
        },
        "time_to_x_start_row": 37,
        "defense_scenario_start_row": 42,
        "il_caveat_row": 49,
    }

    _build_time_to_x_callouts(ws, treasury_meta, defense_target_low_ref)
    _build_defense_scenario_table(ws, defense_target_low_ref)
    _build_il_caveat(ws)

    # ------------------------------------------------------------------
    # Charts
    # ------------------------------------------------------------------
    _create_treasury_composition_chart(ws, treasury_meta)
    _create_depletion_chart(ws, treasury_meta)
    _create_buyback_burn_chart(ws, treasury_meta)

    # ------------------------------------------------------------------
    # Conditional formatting
    # ------------------------------------------------------------------
    _add_cp_balance_formatting(ws, treasury_meta)
    _add_defense_treasury_formatting(ws, treasury_meta)
    _add_net_treasury_heatmap(ws, treasury_meta)

    # ------------------------------------------------------------------
    # Column widths
    # ------------------------------------------------------------------
    for col_letter, width in _COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # ------------------------------------------------------------------
    # Freeze panes at A3 (headers always visible)
    # ------------------------------------------------------------------
    ws.freeze_panes = "A3"

    return treasury_meta
