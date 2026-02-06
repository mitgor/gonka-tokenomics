"""
Gonka Tokenomics - Documentation Tab (Cover Sheet + TOC)

Builds the "Documentation" worksheet at index 0 with:
  - Title and subtitle
  - Version (v1.1) and dynamic date (=TODAY())
  - Disclaimer in red italic
  - Color convention legend with styled sample cells
  - Hyperlinked Table of Contents for all other tabs

Public API:
    build_documentation_tab(wb, all_sheet_names) -> None
"""

from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import quote_sheetname

from generators.styles import TAB_COLOR_DOCS, SECTION_FONT_COLOR


# ---------------------------------------------------------------------------
# Module-level constants
# ---------------------------------------------------------------------------

LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")

TAB_DESCRIPTIONS = {
    "Assumptions": "All adjustable model inputs and scenario selector",
    "Emission Schedule": "Mining reward decay curve and circulating supply",
    "Token Price": "Multi-scenario price trajectories and market cap",
    "Fee Transition": "Fee revenue vs emission crossover analysis",
    "Host Profitability": "GPU host economics and breakeven analysis",
    "Treasury & POL": "Community Pool, POL, buyback-burn, and defense",
    "Dashboard": "Cross-model KPIs, scenario comparison, and summary charts",
}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def build_documentation_tab(wb, all_sheet_names):
    """Build the Documentation (cover/title) sheet at position 0.

    Creates a worksheet named "Documentation" inserted at index 0 and
    populates it with title, version, date, disclaimer, color legend,
    and a hyperlinked Table of Contents.

    Args:
        wb: An openpyxl Workbook instance (styles already registered).
        all_sheet_names: List of sheet names to include in the TOC
            (should NOT include "Documentation" itself).

    Returns:
        None -- the function modifies ``wb`` in-place.
    """
    ws = wb.create_sheet("Documentation", index=0)
    ws.sheet_properties.tabColor = TAB_COLOR_DOCS

    # =================================================================
    # Section 1: Title (rows 1-2)
    # =================================================================
    ws.merge_cells("A1:F1")
    title_cell = ws.cell(row=1, column=1, value="GONKA TOKENOMICS MODEL")
    title_cell.font = Font(
        name="Calibri", size=18, bold=True, color=SECTION_FONT_COLOR,
    )
    title_cell.alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:F2")
    subtitle_cell = ws.cell(
        row=2, column=1,
        value="Gonka Network Token Economics -- Interactive Scenario Model",
    )
    subtitle_cell.alignment = Alignment(horizontal="center")

    # =================================================================
    # Section 2: Version and Date (rows 4-5)
    # =================================================================
    ws.cell(row=4, column=1, value="Version:").font = Font(
        name="Calibri", size=11, bold=True,
    )
    ws.cell(row=4, column=2, value="v1.1")

    ws.cell(row=5, column=1, value="Generated:").font = Font(
        name="Calibri", size=11, bold=True,
    )
    date_cell = ws.cell(row=5, column=2, value="=TODAY()")
    date_cell.number_format = "YYYY-MM-DD"

    # =================================================================
    # Section 3: Disclaimer (row 7)
    # =================================================================
    ws.merge_cells("A7:F7")
    disclaimer_cell = ws.cell(
        row=7, column=1,
        value="For internal decision-making purposes only",
    )
    disclaimer_cell.font = Font(
        name="Calibri", size=11, italic=True, color="FF0000",
    )

    # =================================================================
    # Section 4: Color Convention Legend (rows 9-14)
    # =================================================================
    ws.cell(row=9, column=1, value="COLOR CONVENTIONS").style = "section_header"

    _legend_entries = [
        ("Blue cells = Editable inputs (Assumptions tab)", "input_cell"),
        ("Black text = Calculated formulas", "formula_cell"),
        ("Green text = Cross-tab references", "crossref_cell"),
        ("Blue headers = Column headers", "header"),
        ("Red shading = Danger zones / warnings", "warning_cell"),
    ]
    for i, (description, style_name) in enumerate(_legend_entries):
        row = 10 + i
        ws.cell(row=row, column=1, value=description)
        sample_cell = ws.cell(row=row, column=2, value="Example")
        sample_cell.style = style_name

    # =================================================================
    # Section 5: Table of Contents (rows 16+)
    # =================================================================
    ws.cell(row=16, column=1, value="TABLE OF CONTENTS").style = "section_header"

    toc_row = 17
    for idx, name in enumerate(all_sheet_names, start=1):
        # Column A: numbered tab name with hyperlink
        cell = ws.cell(row=toc_row, column=1, value=f"{idx}. {name}")
        quoted_name = quote_sheetname(name)
        cell.hyperlink = f"#{quoted_name}!A1"
        cell.font = LINK_FONT

        # Column B: brief description
        ws.cell(row=toc_row, column=2, value=TAB_DESCRIPTIONS.get(name, ""))

        toc_row += 1

    # =================================================================
    # Tab formatting
    # =================================================================
    col_widths = {"A": 45, "B": 25, "C": 15, "D": 15, "E": 15, "F": 15}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # No freeze panes on Documentation tab
