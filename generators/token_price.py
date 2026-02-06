"""
Gonka Tokenomics - Token Price Scenarios Model Tab

Builds the "Token Price" worksheet with 32 data rows (matching the Emission
Schedule period structure) and 12 columns:

  A  Period Label           (cross-ref from Emission Schedule)
  B  Conservative ($)       (linear interpolation low->high)
  C  Moderate ($)           (linear interpolation low->high)
  D  Aggressive ($)         (linear interpolation low->high)
  E  Bitfury ($)            (constant Schelling Point)
  F  Active Price ($)       (interpolation using CHOOSE-derived Active Low/High)
  G  Circulating Supply     (cross-ref from Emission Schedule col H)
  H  FDV ($)                (Total Supply * Active Price)
  I  Circ. Market Cap ($)   (Circulating Supply * Active Price)
  J  Gross New Supply       (cross-ref from Emission Schedule col G)
  K  Buyback Burn (GNK)     (fee_revenue * buyback_pct / price / periods_per_year)
  L  Net Supply Change      (Gross - Buyback)

Every formula references the Assumptions tab via param_refs or the Emission
Schedule tab via emission_meta, so changes propagate automatically.

Public API:
    build_token_price_tab(wb, param_refs, emission_meta) -> price_meta dict
"""

from openpyxl.utils import quote_sheetname

from generators.chart_utils import col_to_num
from generators.styles import TAB_COLOR_CALC


# ---------------------------------------------------------------------------
# Column layout
# ---------------------------------------------------------------------------

_HEADERS = [
    "Period",               # A
    "Conservative ($)",     # B
    "Moderate ($)",         # C
    "Aggressive ($)",       # D
    "Bitfury ($)",          # E
    "Active Price ($)",     # F
    "Circulating Supply",   # G
    "FDV ($)",              # H
    "Circ. Market Cap ($)", # I
    "Gross New Supply",     # J
    "Buyback Burn (GNK)",   # K
    "Net Supply Change",    # L
]

_COL_WIDTHS = {
    "A": 12, "B": 16, "C": 14, "D": 16, "E": 12,
    "F": 16, "G": 22, "H": 18, "I": 20, "J": 18,
    "K": 18, "L": 18,
}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def build_token_price_tab(wb, param_refs, emission_meta):
    """Create the 'Token Price' worksheet and populate formulas.

    Args:
        wb: An openpyxl Workbook with styles already registered.
        param_refs: dict mapping parameter names to "Assumptions!$B$N".
                    Must include scenario selector refs (Active Price Low/High).
        emission_meta: dict from build_emission_tab() with sheet coordinates.

    Returns:
        dict: price_meta with sheet coordinates for downstream tabs/charts.
    """
    ws = wb.create_sheet(title="Token Price")
    ws.sheet_properties.tabColor = TAB_COLOR_CALC

    # ------------------------------------------------------------------
    # Resolve param_refs we need
    # ------------------------------------------------------------------
    cons_low = param_refs["Conservative Price Low"]
    cons_high = param_refs["Conservative Price High"]
    mod_low = param_refs["Moderate Price Low"]
    mod_high = param_refs["Moderate Price High"]
    agg_low = param_refs["Aggressive Price Low"]
    agg_high = param_refs["Aggressive Price High"]
    bitfury_ref = param_refs["Bitfury Schelling Point"]
    active_low = param_refs["Active Price Low"]
    active_high = param_refs["Active Price High"]
    total_supply_ref = param_refs["Total Supply"]
    fee_rev_ref = param_refs["Assumed Annual Fee Revenue"]
    buyback_ref = param_refs["Buyback-Burn"]

    # Emission Schedule cross-sheet references
    es_sheet = quote_sheetname(emission_meta["sheet_name"])
    circ_col = emission_meta["cols"]["cumulative_circulating"]   # "H"
    total_col = emission_meta["cols"]["total_new"]               # "G"
    period_col = emission_meta["cols"]["period_label"]           # "A"

    # ------------------------------------------------------------------
    # Row 1: Merged title
    # ------------------------------------------------------------------
    ws.merge_cells("A1:L1")
    title_cell = ws.cell(row=1, column=1, value="TOKEN PRICE SCENARIOS MODEL")
    title_cell.style = "section_header"

    # ------------------------------------------------------------------
    # Row 2: Column headers
    # ------------------------------------------------------------------
    for col_idx, header_text in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header_text)
        cell.style = "header"

    # ------------------------------------------------------------------
    # Rows 3-34: Data rows (32 periods)
    # ------------------------------------------------------------------
    data_start_row = 3
    num_periods = 32
    last_index = num_periods - 1  # 31 (for interpolation denominator)

    for i in range(num_periods):
        row = data_start_row + i
        es_row = emission_meta["data_start_row"] + i

        # Determine periods_per_year for buyback calculation
        # i=0..23 are monthly (12 per year), i=24..31 are annual (1 per year)
        periods_per_year = 12 if i < 24 else 1

        # A: Period Label (cross-ref from Emission Schedule)
        a_cell = ws.cell(
            row=row, column=1,
            value=f"={es_sheet}!{period_col}{es_row}",
        )
        a_cell.style = "crossref_cell"

        # B: Conservative Price (linear interpolation)
        b_cell = ws.cell(
            row=row, column=2,
            value=f"={cons_low}+({i}/{last_index})*({cons_high}-{cons_low})",
        )
        b_cell.style = "currency_precise"

        # C: Moderate Price (linear interpolation)
        c_cell = ws.cell(
            row=row, column=3,
            value=f"={mod_low}+({i}/{last_index})*({mod_high}-{mod_low})",
        )
        c_cell.style = "currency_precise"

        # D: Aggressive Price (linear interpolation)
        d_cell = ws.cell(
            row=row, column=4,
            value=f"={agg_low}+({i}/{last_index})*({agg_high}-{agg_low})",
        )
        d_cell.style = "currency_precise"

        # E: Bitfury Schelling Point (constant flat line)
        e_cell = ws.cell(
            row=row, column=5,
            value=f"={bitfury_ref}",
        )
        e_cell.style = "currency_precise"

        # F: Active Price (interpolation using CHOOSE-derived Low/High)
        f_cell = ws.cell(
            row=row, column=6,
            value=f"={active_low}+({i}/{last_index})*({active_high}-{active_low})",
        )
        f_cell.style = "currency_precise"

        # G: Circulating Supply (cross-ref from Emission Schedule col H)
        g_cell = ws.cell(
            row=row, column=7,
            value=f"={es_sheet}!{circ_col}{es_row}",
        )
        g_cell.style = "crossref_cell"
        g_cell.number_format = "#,##0"

        # H: FDV = Total Supply * Active Price
        h_cell = ws.cell(
            row=row, column=8,
            value=f"={total_supply_ref}*F{row}",
        )
        h_cell.style = "currency"

        # I: Circulating Market Cap = Circulating Supply * Active Price
        i_cell = ws.cell(
            row=row, column=9,
            value=f"=G{row}*F{row}",
        )
        i_cell.style = "currency"

        # J: Gross New Supply (cross-ref from Emission Schedule col G)
        j_cell = ws.cell(
            row=row, column=10,
            value=f"={es_sheet}!{total_col}{es_row}",
        )
        j_cell.style = "crossref_cell"
        j_cell.number_format = "#,##0"

        # K: Buyback Burn (GNK) = fee_rev * buyback_pct / price / periods_per_year
        k_cell = ws.cell(
            row=row, column=11,
            value=f"={fee_rev_ref}*{buyback_ref}/F{row}/{periods_per_year}",
        )
        k_cell.style = "tokens"

        # L: Net Supply Change = Gross New Supply - Buyback Burn
        l_cell = ws.cell(
            row=row, column=12,
            value=f"=J{row}-K{row}",
        )
        l_cell.style = "tokens"

    # ------------------------------------------------------------------
    # Column widths
    # ------------------------------------------------------------------
    for col_letter, width in _COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # ------------------------------------------------------------------
    # Freeze panes at A3 (headers always visible)
    # ------------------------------------------------------------------
    ws.freeze_panes = "A3"

    # ------------------------------------------------------------------
    # Build price_meta for downstream tabs and charts
    # ------------------------------------------------------------------
    price_meta = {
        "sheet_name": "Token Price",
        "header_row": 2,
        "data_start_row": 3,
        "data_end_row": 34,
        "cols": {
            "period_label": "A",
            "conservative_price": "B",
            "moderate_price": "C",
            "aggressive_price": "D",
            "bitfury_price": "E",
            "active_price": "F",
            "circulating_supply": "G",
            "fdv": "H",
            "circ_market_cap": "I",
            "gross_new_supply": "J",
            "buyback_burn": "K",
            "net_supply_change": "L",
        },
    }

    return price_meta
