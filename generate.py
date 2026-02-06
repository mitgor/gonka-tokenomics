#!/usr/bin/env python3
"""
Gonka Tokenomics Model Generator

Generates Excel workbooks (.xlsx) for Gonka Network tokenomics modeling.
All parameters from v1.0 research are written to an Assumptions tab that
drives all downstream calculations via Excel formulas.

Usage:
    python generate.py              # Generate all workbooks
    python generate.py --test       # Generate minimal test workbook (Assumptions only)

Output directory: ./output/
"""
import sys
from pathlib import Path


def main():
    # Ensure output directory exists
    output_dir = Path("output")
    output_dir.mkdir(exist_ok=True)

    # Parse args
    test_mode = "--test" in sys.argv

    if test_mode:
        generate_test()
    else:
        generate_all()


def generate_test():
    """Generate a minimal test workbook with only the Assumptions tab."""
    from generators.workbook_base import create_workbook

    wb, param_refs = create_workbook()

    output_path = Path("output") / "gonka_test_assumptions.xlsx"
    wb.save(str(output_path))
    print(f"Generated: {output_path}")
    print(f"  Assumptions tab: {len(param_refs)} parameters")


def _add_back_to_doc_links(wb):
    """Add 'Back to Documentation' hyperlink on every non-Documentation tab.

    IMPORTANT: Dashboard is EXCLUDED here. Plan 02 builds Dashboard content
    with a merged title at A1:E1, so it adds its own '<< Documentation' link
    at F1 after building content. Including Dashboard here would place a link
    at A1 that gets overwritten by the title merge.
    """
    from openpyxl.styles import Font
    link_font = Font(name="Calibri", size=11, color="0563C1", underline="single")

    BACK_LINK_COL_ROW = {
        "Assumptions": (6, 1),          # F1 (after A1:E1 merge)
        "Emission Schedule": (11, 2),   # K2 (K1 is chart anchor; use row 2)
        "Token Price": (13, 1),         # M1 (after A1:L1 merge, before N1 chart)
        "Fee Transition": (15, 1),      # O1 (after A1:N1 merge, before P1 chart)
        "Host Profitability": (14, 1),  # N1 (after A1:M1 merge, before O1 chart)
        "Treasury & POL": (15, 1),      # O1 (after A1:N1 merge, before P1 chart)
        # Dashboard excluded -- Plan 02 adds link at F1 after building content
    }

    for ws in wb.worksheets:
        if ws.title == "Documentation" or ws.title == "Dashboard":
            continue
        if ws.title in BACK_LINK_COL_ROW:
            col, row = BACK_LINK_COL_ROW[ws.title]
        else:
            col, row = 1, 1
        cell = ws.cell(row=row, column=col)
        cell.value = "<< Documentation"
        cell.hyperlink = "#Documentation!A1"
        cell.font = link_font


def generate_all():
    """Generate the master workbook with all model tabs and charts."""
    from generators.workbook_base import create_workbook
    from generators.emission import build_emission_tab
    from generators.token_price import build_token_price_tab
    from generators.chart_utils import fix_chart_rendering
    from generators.fee_transition import build_fee_transition_tab
    from generators.host_profit import build_host_profit_tab
    from generators.treasury import build_treasury_tab
    from generators.documentation import build_documentation_tab
    from generators.dashboard import build_dashboard_tab

    wb, param_refs = create_workbook()

    # Phase 2: Emission Schedule Model
    emission_meta = build_emission_tab(wb, param_refs)

    # Phase 3: Token Price Scenarios Model
    price_meta = build_token_price_tab(wb, param_refs, emission_meta)

    # Phase 4: Fee Transition Crossover Model
    fee_meta = build_fee_transition_tab(wb, param_refs, emission_meta, price_meta)

    # Phase 5: Host Profitability Model
    host_meta = build_host_profit_tab(wb, param_refs, emission_meta, price_meta, fee_meta)

    # Phase 6: Treasury & POL Simulation
    treasury_meta = build_treasury_tab(wb, param_refs, emission_meta, price_meta, fee_meta)

    # Phase 7: Dashboard tab with KPIs, scenario matrix, and charts
    ws_dashboard = wb.create_sheet("Dashboard")
    ws_dashboard.sheet_properties.tabColor = "ED7D31"  # TAB_COLOR_DASHBOARD
    dashboard_meta = build_dashboard_tab(wb, param_refs, emission_meta, price_meta, fee_meta, host_meta, treasury_meta)

    # Phase 7: Documentation tab (inserted at position 0)
    EXPECTED_TABS = [
        "Documentation", "Assumptions", "Emission Schedule", "Token Price",
        "Fee Transition", "Host Profitability", "Treasury & POL", "Dashboard",
    ]
    build_documentation_tab(wb, EXPECTED_TABS[1:])  # Pass all tabs except Documentation itself

    # Assert 8-tab order
    assert wb.sheetnames == EXPECTED_TABS, f"Tab order mismatch: {wb.sheetnames} != {EXPECTED_TABS}"

    # Add "Back to Documentation" links on every tab except Documentation and Dashboard
    _add_back_to_doc_links(wb)

    output_path = Path("output") / "gonka_master_model.xlsx"
    wb.save(str(output_path))
    fix_chart_rendering(str(output_path))
    print(f"Generated: {output_path}")
    print(f"  Documentation tab: cover sheet + TOC with {len(EXPECTED_TABS) - 1} navigation links")
    print(f"  Assumptions tab: {len(param_refs)} parameters")
    print(f"  Emission Schedule tab: {emission_meta['data_end_row'] - emission_meta['data_start_row'] + 1} periods, 3 charts")
    print(f"  Token Price tab: {price_meta['data_end_row'] - price_meta['data_start_row'] + 1} periods, 2 charts")
    print(f"  Fee Transition tab: {fee_meta['data_end_row'] - fee_meta['data_start_row'] + 1} periods, 14 columns, 2 matrices")
    print(f"  Host Profitability tab: {host_meta['data_end_row'] - host_meta['data_start_row'] + 1} periods, 13 columns, sensitivity matrix")
    print(f"  Treasury & POL tab: {treasury_meta['data_end_row'] - treasury_meta['data_start_row'] + 1} periods, 14 columns")
    print(f"  Dashboard tab: {dashboard_meta['kpi_end_row'] - dashboard_meta['kpi_start_row'] + 1} KPIs, scenario matrix, 3 charts")


if __name__ == "__main__":
    main()
